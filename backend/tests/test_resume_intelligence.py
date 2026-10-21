from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from app.main import app
from app.nlp.resume_cleaner import clean_resume_text
from app.nlp.resume_extractor import (
    calculate_years_experience,
    detect_sections,
    extract_resume_profile,
)
from app.parsers.base import ResumeParserError
from app.parsers.docx_parser import DOCXResumeParser
from app.parsers.pdf_parser import PDFResumeParser
from app.schemas.resume import ExperienceItem
from app.utils.file_validation import (
    DOCX_MIME,
    MAX_RESUME_BYTES,
    UploadValidationError,
    sanitize_filename,
    validate_resume_upload,
)


def make_docx() -> bytes:
    document = Document()
    document.add_heading('Alex Kumar', level=1)
    document.add_paragraph('alex.kumar@example.com')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Python'
    table.cell(0, 1).text = 'PostgreSQL'
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_text_pdf(pages: list[str]) -> bytes:
    """Create a tiny standards-compliant synthetic text PDF without external fixtures."""
    page_count = len(pages)
    font_id = 3 + page_count
    content_start = font_id + 1
    objects: list[bytes] = []
    objects.append(b'<< /Type /Catalog /Pages 2 0 R >>')
    kids = ' '.join(f'{3 + i} 0 R' for i in range(page_count))
    objects.append(f'<< /Type /Pages /Kids [{kids}] /Count {page_count} >>'.encode())
    for index in range(page_count):
        objects.append(
            f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
            f'/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_start + index} 0 R >>'.encode()
        )
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    for text in pages:
        escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        stream = f'BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET'.encode('ascii')
        objects.append(f'<< /Length {len(stream)} >>\nstream\n'.encode() + stream + b'\nendstream')

    output = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f'{index} 0 obj\n'.encode() + body + b'\nendobj\n')
    xref_offset = len(output)
    output.extend(f'xref\n0 {len(objects) + 1}\n'.encode())
    output.extend(b'0000000000 65535 f \n')
    for offset in offsets[1:]:
        output.extend(f'{offset:010d} 00000 n \n'.encode())
    output.extend(f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n'.encode())
    return bytes(output)


def test_file_validation_accepts_pdf_docx_and_sanitizes_filename():
    pdf = make_text_pdf(['Synthetic resume'])
    validated_pdf = validate_resume_upload('../../Alex Resume.PDF', pdf)
    assert validated_pdf.file_name == 'Alex Resume.PDF'
    assert validated_pdf.file_type == 'application/pdf'
    validated_docx = validate_resume_upload('Alex.docx', make_docx())
    assert validated_docx.file_type == DOCX_MIME
    assert sanitize_filename('..\\private\\bad:name.docx') == 'bad_name.docx'


@pytest.mark.parametrize(('filename', 'data'), [
    ('virus.exe', b'MZnot-a-resume'),
    ('photo.jpg', b'\xff\xd8\xff'),
    ('archive.zip', b'PK\x03\x04fake'),
    ('report.pdf', b'not a PDF'),
    ('report.docx', b'not a DOCX'),
])
def test_file_validation_rejects_unsupported_or_mismatched_content(filename, data):
    with pytest.raises(UploadValidationError):
        validate_resume_upload(filename, data)


def test_file_validation_rejects_empty_and_oversized_files():
    with pytest.raises(UploadValidationError, match='empty'):
        validate_resume_upload('empty.pdf', b'')
    with pytest.raises(UploadValidationError, match='10 MB'):
        validate_resume_upload('large.pdf', b'%PDF-' + b'x' * MAX_RESUME_BYTES)


def test_docx_validation_rejects_corrupt_archive_with_docx_extension():
    # A valid zip lacking the required OOXML parts is not a DOCX.
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        archive.writestr('readme.txt', 'not a docx')
    with pytest.raises(UploadValidationError, match='valid DOCX'):
        validate_resume_upload('fake.docx', output.getvalue())


def test_pdf_parser_extracts_text_across_multiple_pages():
    parsed = PDFResumeParser().parse(make_text_pdf(['Alex Kumar', 'Python and SQL']))
    assert parsed.page_count == 2
    assert 'Alex Kumar' in parsed.text
    assert 'Python and SQL' in parsed.text


def test_pdf_parser_rejects_invalid_and_image_only_pdfs():
    with pytest.raises(ResumeParserError, match='valid, unencrypted'):
        PDFResumeParser().parse(b'%PDF-not-readable')
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ResumeParserError, match='image-only PDF'):
        PDFResumeParser().parse(output.getvalue())


def test_pdf_parser_reports_password_protected_document():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt('secret')
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ResumeParserError, match='password-protected'):
        PDFResumeParser().parse(output.getvalue())


def test_docx_parser_extracts_paragraphs_and_table_text():
    parsed = DOCXResumeParser().parse(make_docx())
    assert 'Alex Kumar' in parsed.text
    assert 'Python | PostgreSQL' in parsed.text
    assert parsed.metadata['paragraph_count'] >= 2


def test_docx_parser_handles_corrupt_content():
    with pytest.raises(ResumeParserError, match='DOCX file could not be read'):
        DOCXResumeParser().parse(b'PK\x03\x04corrupt')


def test_cleaner_normalizes_noise_but_preserves_technical_punctuation():
    cleaned = clean_resume_text('  C++   C#\r\n\r\n\u00a0.NET\x00\nNode.js\nSQL')
    assert 'C++ C#' in cleaned
    assert '.NET' in cleaned
    assert 'Node.js' in cleaned
    assert '\x00' not in cleaned
    assert '\n\n\n' not in cleaned
    wrapped = clean_resume_text('A sufficiently long resume sentence that wrapped across the page\nonto the next extracted line.')
    assert wrapped == 'A sufficiently long resume sentence that wrapped across the page onto the next extracted line.'


def test_cleaner_removes_repeated_short_page_headers():
    cleaned = clean_resume_text('Alex Kumar\nExperience\nAlex Kumar\nProject work\nAlex Kumar\nEducation')
    assert cleaned.count('Alex Kumar') == 0
    assert 'Project work' in cleaned


def test_sections_and_structured_resume_extraction():
    fixture_path = Path(__file__).parent / 'fixtures' / 'sample_resume.txt'
    fixture = open(fixture_path, encoding='utf-8').read()
    cleaned = clean_resume_text(fixture)
    sections = detect_sections(cleaned)
    assert {'summary', 'education', 'experience', 'skills', 'projects', 'certifications', 'achievements', 'languages'} <= sections.keys()

    profile, _ = extract_resume_profile(cleaned)
    assert profile.full_name == 'Alex Kumar'
    assert profile.email == 'alex.kumar@example.com'
    assert profile.phone is not None and '90000' in profile.phone
    assert profile.location == 'Chennai, India'
    assert profile.summary and 'Data Science student' in profile.summary
    assert any(item.institution == 'Example University, Chennai' for item in profile.education)
    assert profile.experience and profile.experience[0].job_title == 'Software Engineer'
    assert profile.experience[0].company == 'Example Tech'
    assert profile.experience[0].description and 'internal tools' in profile.experience[0].description
    assert profile.years_experience == 1.0
    assert 'C++' in profile.skills.programming_languages
    assert 'Machine Learning' in profile.skills.technical
    assert profile.projects and profile.projects[0].project_name == 'Resume Screening System'
    assert profile.certifications[0].issuer == 'Example Institute'
    assert any(language.name == 'Tamil' and language.proficiency == 'Native' for language in profile.languages)
    assert any('Hackathon' in achievement.title for achievement in profile.achievements)
    assert profile.sections.education is True


def test_experience_calculation_merges_overlapping_ranges_and_ignores_bad_dates():
    entries = [
        ExperienceItem(start_date='2020', end_date='2022'),
        ExperienceItem(start_date='2021', end_date='2023'),
        ExperienceItem(start_date='2025', end_date='2020'),
        ExperienceItem(start_date=None, end_date='Present'),
    ]
    assert calculate_years_experience(entries) == 3.0
    assert calculate_years_experience([ExperienceItem(job_title='Undated')]) is None


def test_resume_upload_openapi_and_authentication_requirement():
    client = TestClient(app)
    response = client.post('/api/resumes/upload')
    assert response.status_code == 401
    schema = client.get('/openapi.json').json()
    assert '/api/resumes/upload' in schema['paths']
    assert 'multipart/form-data' in schema['paths']['/api/resumes/upload']['post']['requestBody']['content']


def test_cors_preflight_allows_accept_header():
    client = TestClient(app)
    response = client.options(
        '/api/health',
        headers={
            'Origin': 'http://localhost:5173',
            'Access-Control-Request-Method': 'GET',
            'Access-Control-Request-Headers': 'authorization, accept, content-type',
        },
    )
    assert response.status_code == 200
    allow_headers = response.headers.get('access-control-allow-headers', '').lower()
    assert 'accept' in allow_headers


def test_demo_token_fallback_in_development():
    client = TestClient(app)
    response = client.get(
        '/api/career-assistant/sessions',
        headers={'Authorization': 'Bearer demo-token-12345'},
    )
    # The dependency accepts the demo token and doesn't reject with 401
    assert response.status_code != 401
