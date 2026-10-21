"""DOCX text extraction preserving paragraph/table order."""

from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsers.base import ParsedDocument, ResumeParserError


class DOCXResumeParser:
    def parse(self, file_bytes: bytes) -> ParsedDocument:
        try:
            document = Document(BytesIO(file_bytes))
            parts: list[str] = []
            for child in document.element.body.iterchildren():
                if isinstance(child, CT_P):
                    paragraph = Paragraph(child, document)
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
                elif isinstance(child, CT_Tbl):
                    table = Table(child, document)
                    for row in table.rows:
                        cells = [" ".join(cell.text.split()) for cell in row.cells]
                        line = " | ".join(cell for cell in cells if cell)
                        if line:
                            parts.append(line)
            text = "\n".join(parts)
            if not text.strip():
                raise ResumeParserError("No readable text was found in this DOCX file.")
            return ParsedDocument(text=text, page_count=None, metadata={"paragraph_count": len(document.paragraphs)})
        except ResumeParserError:
            raise
        except Exception as error:
            raise ResumeParserError("This DOCX file could not be read. Save a fresh .docx copy and try again.") from error
